"""Generates documentation/Manuscript_Update_Brief.docx from Manuscript_Update_Brief.md."""
from docx import Document
from docx.shared import Pt
import re

SRC = "documentation/Manuscript_Update_Brief.md"
OUT = "documentation/Manuscript_Update_Brief.docx"

doc = Document()
doc.styles["Normal"].font.size = Pt(11)

def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            if j < len(t.rows[i].cells):
                t.rows[i].cells[j].text = val
                for p in t.rows[i].cells[j].paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(10)
                        if i == 0:
                            run.font.bold = True

with open(SRC, encoding="utf-8") as f:
    lines = f.read().splitlines()

i = 0
while i < len(lines):
    line = lines[i]
    if line.startswith("```"):
        i += 1
        code = []
        while i < len(lines) and not lines[i].startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        run = p.add_run("\n".join(code))
        run.font.name = "Consolas"; run.font.size = Pt(9)
    elif line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").replace(" ", "")) <= set("-:"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            rows.append(lines[i]); i += 1
        add_table(rows)
        continue
    elif line.startswith("# "):
        doc.add_heading(line[2:], level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:], level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:], level=2)
    elif re.match(r"^\d+\.", line.strip()):
        doc.add_paragraph(line.strip()[3:], style="List Number")
    elif line.strip().startswith(("- ", "* ")):
        text = line.strip()[2:]
        if text.startswith(("- ", "* ")):  # nested bullet
            doc.add_paragraph(text[2:], style="List Bullet 2")
        else:
            doc.add_paragraph(text, style="List Bullet")
    elif line.strip():
        doc.add_paragraph(line.replace("**", ""))
    i += 1

doc.save(OUT)
print("saved", OUT)
