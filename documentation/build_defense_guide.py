"""Regenerates documentation/System_Defense_Guide.docx from System_Defense_Guide.md.

The Markdown file is the single source of truth for the defence guide; the DOCX is a
build artifact. Run this from the repository root after editing the Markdown:

    python documentation/build_defense_guide.py

(Any Python 3 with `python-docx` installed works; the repository's `ai-engine/.venv`
does not carry it, so use the system interpreter.)

Structure preserved from the Markdown: the title, heading levels 1-3, bullet and
numbered lists (including one nesting level), pipe tables, block quotes, fenced code
blocks (monospaced, used for the Mermaid sources and shell commands), and inline
`code` / **bold** / *italic* runs.
"""

import re

from docx import Document
from docx.shared import Pt

SRC = "documentation/System_Defense_Guide.md"
OUT = "documentation/System_Defense_Guide.docx"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_runs(paragraph, text):
    """Add text to a paragraph, turning Markdown inline markers into real runs."""
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


def is_table_start(lines, i):
    """A pipe row followed by a separator row is a Markdown table."""
    if not lines[i].lstrip().startswith("|") or i + 1 >= len(lines):
        return False
    sep = lines[i + 1].replace("|", "").replace(" ", "")
    return bool(sep) and set(sep) <= set("-:")


def add_table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(set(c) <= set("-: ") for c in r)]
    if not cells:
        return
    table = doc.add_table(rows=len(cells), cols=len(cells[0]))
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            if j >= len(table.rows[i].cells):
                continue
            cell = table.rows[i].cells[j]
            cell.text = ""
            add_runs(cell.paragraphs[0], value)
            for para in cell.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10)
                    if i == 0:
                        run.bold = True


def build():
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.core_properties.title = "IT Faculty Scheduling System — Design Defense Guide"
    doc.core_properties.comments = f"Generated from {SRC} by documentation/build_defense_guide.py"

    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            paragraph = doc.add_paragraph()
            run = paragraph.add_run("\n".join(code))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue

        if is_table_start(lines, i):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            heading = doc.add_heading("", level=min(level - 1, 3))
            add_runs(heading, text)
        elif stripped in ("---", "***", "___"):
            pass  # horizontal rule: the heading hierarchy already separates sections
        elif stripped.startswith(">"):
            paragraph = doc.add_paragraph(style="Intense Quote")
            add_runs(paragraph, stripped.lstrip("> ").strip())
        elif re.match(r"^\d+\.\s", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_runs(paragraph, re.sub(r"^\d+\.\s", "", stripped))
        elif stripped.startswith(("- ", "* ")):
            indent = len(line) - len(line.lstrip())
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            add_runs(paragraph, stripped[2:])
        elif stripped:
            paragraph = doc.add_paragraph()
            add_runs(paragraph, stripped)
        i += 1

    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
